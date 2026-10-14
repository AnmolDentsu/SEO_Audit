from __future__ import annotations

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app import models
from app.database import Base
from app.services.sitemap_export import build_sitemap_workbook


def _sheet_rows(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows or rows[0] == ("No data",):
        return []
    headers = list(rows[0])
    return [dict(zip(headers, row)) for row in rows[1:]]


def test_partial_export_rebuilds_large_views_from_db_not_analysis_json():
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    Session = sessionmaker(bind=engine)
    db = Session()
    try:
        project = models.Project(name="Export Project")
        db.add(project)
        db.flush()

        client = models.SitemapSource(
            project_id=project.id,
            root_url="https://client.example/sitemap.xml",
            role="CLIENT",
            label="Client",
            enabled=True,
        )
        competitor = models.SitemapSource(
            project_id=project.id,
            root_url="https://competitor.example/sitemap.xml",
            role="COMPETITOR",
            label="Competitor",
            enabled=True,
        )
        db.add_all([client, competitor])
        db.flush()

        run = models.SitemapRun(
            project_id=project.id,
            run_number=15,
            status="partial",
            phase="partial",
            config_json={
                "_input_snapshot": {
                    "sources": [
                        {"id": client.id, "root_url": client.root_url, "role": "CLIENT", "label": "Client", "enabled": True},
                        {"id": competitor.id, "root_url": competitor.root_url, "role": "COMPETITOR", "label": "Competitor", "enabled": True},
                    ]
                },
                "_background_recovery": {"state": "cooling", "next_attempt_at": "2026-10-08T12:00:00+00:00"},
            },
            analysis_json={
                "overview": [{"Metric": "Persisted", "Value": 3}],
                "source_discovery": [],
                "site_profiles": [],
                "url_type_summary": [],
                "comparison_status": {"ready": False, "message": "Partial evidence"},
                # Deliberately NO large taxonomy/coverage/head-to-head arrays.
                "sitemap_failures": [
                    {
                        "source_id": competitor.id,
                        "site": "Competitor",
                        "role": "COMPETITOR",
                        "root_sitemap": competitor.root_url,
                        "parent_sitemap": competitor.root_url,
                        "url": "https://competitor.example/empty-child.xml",
                        "stage": "sitemap_discovery",
                        "severity": "warning",
                        "code": "EMPTY_CHILD_SITEMAP",
                        "http_status": 200,
                        "verification_attempts": 1,
                        "empty_kind": "empty_body",
                        "error": "Confirmed empty child",
                    }
                ],
            },
        )
        db.add(run)
        db.flush()

        rows = [
            models.SitemapUrl(
                sitemap_run_id=run.id, source_id=client.id,
                root_sitemap=client.root_url, child_sitemap=client.root_url,
                original_url="https://client.example/loans", normalized_url="https://client.example/loans",
                final_url="https://client.example/loans", http_status=200, fetch_method="http",
                broad_page_type="PRODUCT", detailed_page_type="PRODUCT", type_confidence=.9,
                classification_method="deterministic", classification_evidence_json={"terminal_state": "usable", "taxonomy_eligible": True},
                category="Loans", topic="Personal Loan", subtopic="Eligibility", topic_confidence=.9,
                topic_method="deterministic", title="Personal Loan", h1_json=["Personal Loan"], h2_json=[],
                schema_types_json=[], breadcrumbs_json=[], word_count=500, author_present=False, cta_json=[], error=None,
            ),
            models.SitemapUrl(
                sitemap_run_id=run.id, source_id=competitor.id,
                root_sitemap=competitor.root_url, child_sitemap=competitor.root_url,
                original_url="https://competitor.example/loans", normalized_url="https://competitor.example/loans",
                final_url="https://competitor.example/loans", http_status=200, fetch_method="http",
                broad_page_type="CONTENT", detailed_page_type="CONTENT", type_confidence=.9,
                classification_method="deterministic", classification_evidence_json={"terminal_state": "usable", "taxonomy_eligible": True},
                category="Loans", topic="Personal Loan", subtopic="Documents", topic_confidence=.9,
                topic_method="deterministic", title="Loan Guide", h1_json=["Loan Guide"], h2_json=[],
                schema_types_json=[], breadcrumbs_json=[], word_count=800, author_present=False, cta_json=[], error=None,
            ),
            models.SitemapUrl(
                sitemap_run_id=run.id, source_id=competitor.id,
                root_sitemap=competitor.root_url, child_sitemap=competitor.root_url,
                original_url="https://competitor.example/blocked", normalized_url="https://competitor.example/blocked",
                final_url="https://competitor.example/blocked", http_status=403, fetch_method="curl-cffi",
                broad_page_type="OTHER", detailed_page_type="OTHER", type_confidence=0,
                classification_method="fetch-error", classification_evidence_json={"terminal_state": "fetch_failed", "taxonomy_eligible": False},
                category="Other / Unclassified", topic="Other / Unclassified", topic_confidence=0,
                topic_method="not-classified", h1_json=[], h2_json=[], schema_types_json=[], breadcrumbs_json=[],
                word_count=0, author_present=False, cta_json=[], error="HTTP 403",
            ),
        ]
        db.add_all(rows)
        db.commit()

        out = build_sitemap_workbook(db, run.id)
        wb = load_workbook(out, read_only=True, data_only=True)

        assert "Run Status" in wb.sheetnames
        assert "Topic Taxonomy" in wb.sheetnames
        assert "Dynamic Topic Clusters" in wb.sheetnames
        assert "Topic Coverage Matrix" in wb.sheetnames
        assert "Head-to-Head" in wb.sheetnames
        assert "Missing Categories" in wb.sheetnames
        assert "Sitemap Discovery Audit" in wb.sheetnames
        assert "Sitemap Failures" in wb.sheetnames

        status = _sheet_rows(wb["Run Status"])[0]
        assert status["Analysis State"] == "PARTIAL"
        assert status["Persisted URLs"] == 3
        assert status["Background Recovery State"] == "cooling"

        taxonomy = _sheet_rows(wb["Topic Taxonomy"])
        assert any(x["Topic"] == "Personal Loan" for x in taxonomy)

        coverage = _sheet_rows(wb["Topic Coverage Matrix"])
        assert any(x["Topic"] == "Personal Loan" for x in coverage)

        head = _sheet_rows(wb["Head-to-Head"])
        assert head

        discovery_audit = _sheet_rows(wb["Sitemap Discovery Audit"])
        assert any(x["Code"] == "EMPTY_CHILD_SITEMAP" for x in discovery_audit)

        failures = _sheet_rows(wb["Sitemap Failures"])
        assert any(x["Code"] == "EMPTY_CHILD_SITEMAP" for x in failures)
        assert any(x["HTTP Status"] == 403 for x in failures)
    finally:
        db.close()


def test_paged_export_exhausts_every_result_page():
    from app.services.sitemap_export import _paged_view_rows

    values = list(range(1203))

    def fake_view(db, run_id, page=1, page_size=500):
        start = (page - 1) * page_size
        chunk = values[start:start + page_size]
        return {
            "page": page,
            "page_size": page_size,
            "total": len(values),
            "items": [{"value": x} for x in chunk],
        }

    rows = list(_paged_view_rows(fake_view, None, 15, page_size=500))
    assert len(rows) == 1203
    assert rows[0]["value"] == 0
    assert rows[-1]["value"] == 1202
