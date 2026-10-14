from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Iterable, Iterator

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill

from app.database import SessionLocal
from app import models
from app.services.sitemap_result_views import (
    cluster_audit_page,
    failure_page,
    head_to_head_page,
    missing_categories_page,
    topic_coverage_page,
    topic_taxonomy_page,
)


EXCEL_CELL_LIMIT = 32767
PAGE_SIZE = 500

REBUILD_SHEETS = {
    "Topic Taxonomy",
    "Dynamic Topic Clusters",
    "Topic Coverage Matrix",
    "Head-to-Head",
    "Missing Categories",
    "Sitemap Failures",
}


def clean_value(value):
    if value is None:
        return None

    if isinstance(value, (int, float, bool)):
        return value

    if isinstance(value, dict):
        value = json.dumps(
            value,
            ensure_ascii=False,
            separators=(",", ":"),
        )

    elif isinstance(value, (list, tuple, set)):
        value = " | ".join(str(x) for x in value)

    else:
        value = str(value)

    value = ILLEGAL_CHARACTERS_RE.sub("", value)

    value = re.sub(
        r"[\x00-\x08\x0B\x0C\x0E-\x1F\x7F]",
        "",
        value,
    )

    return value[:EXCEL_CELL_LIMIT]


def make_header_cell(ws, value):
    cell = WriteOnlyCell(
        ws,
        value=clean_value(value),
    )

    cell.font = Font(
        bold=True,
        color="FFFFFF",
    )

    cell.fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    return cell


def append_dict_sheet(
    wb: Workbook,
    title: str,
    rows: Iterable[dict],
):
    ws = wb.create_sheet(title[:31])

    iterator = iter(rows)

    first = next(
        iterator,
        None,
    )

    if first is None:
        ws.append(["No data"])
        return

    headers = list(first.keys())

    ws.append([
        make_header_cell(ws, header)
        for header in headers
    ])

    def emit(row):
        ws.append([
            clean_value(row.get(header))
            for header in headers
        ])

    emit(first)

    for row in iterator:
        emit(row)


def copy_existing_sheet(
    source_ws,
    target_wb: Workbook,
):
    target_ws = target_wb.create_sheet(
        source_ws.title[:31]
    )

    for row in source_ws.iter_rows(
        values_only=True
    ):
        target_ws.append([
            clean_value(value)
            for value in row
        ])


def paged_rows(
    fn,
    db,
    run_id: int,
) -> Iterator[dict]:

    page = 1

    while True:

        result = fn(
            db,
            run_id,
            page=page,
            page_size=PAGE_SIZE,
        )

        items = list(
            result.get("items") or []
        )

        for item in items:
            yield item

        total = int(
            result.get("total") or 0
        )

        if not items:
            break

        if page * PAGE_SIZE >= total:
            break

        page += 1


def failure_rows(
    db,
    run_id: int,
) -> Iterator[dict]:

    page = 1
    supplemental_written = False

    while True:

        result = failure_page(
            db,
            run_id,
            page=page,
            page_size=PAGE_SIZE,
        )

        items = list(
            result.get("items") or []
        )

        for item in items:
            yield {
                "Site": item.get("site"),
                "Role": item.get("role"),
                "Stage": item.get("stage"),
                "Host": item.get("host"),
                "URL": item.get("url"),
                "HTTP Status": item.get("http_status"),
                "Retry Count": item.get("retry_count"),
                "Terminal State": item.get("terminal_state"),
                "Severity": item.get("severity"),
                "Code": item.get("code"),
                "Fetch Method": item.get("fetch_method"),
                "Error": item.get("error"),
            }

        if not supplemental_written:

            for item in (
                result.get("supplemental")
                or []
            ):

                yield {
                    "Site": item.get("site"),
                    "Role": item.get("role"),
                    "Stage": item.get("stage"),
                    "Host": item.get("host"),
                    "URL": item.get("url"),
                    "HTTP Status": item.get("http_status"),
                    "Retry Count": item.get("retry_count"),
                    "Terminal State": item.get("terminal_state"),
                    "Severity": item.get("severity"),
                    "Code": item.get("code"),
                    "Fetch Method": item.get("fetch_method"),
                    "Error": item.get("error"),
                }

            supplemental_written = True

        total = int(
            result.get("total") or 0
        )

        if not items:
            break

        if page * PAGE_SIZE >= total:
            break

        page += 1


def rebuild_sheet(
    wb,
    sheet_name,
    db,
    run_id,
):

    print(
        f"Rebuilding: {sheet_name}"
    )

    if sheet_name == "Topic Taxonomy":

        append_dict_sheet(
            wb,
            sheet_name,
            paged_rows(
                topic_taxonomy_page,
                db,
                run_id,
            ),
        )

    elif sheet_name == "Dynamic Topic Clusters":

        append_dict_sheet(
            wb,
            sheet_name,
            paged_rows(
                cluster_audit_page,
                db,
                run_id,
            ),
        )

    elif sheet_name == "Topic Coverage Matrix":

        append_dict_sheet(
            wb,
            sheet_name,
            paged_rows(
                topic_coverage_page,
                db,
                run_id,
            ),
        )

    elif sheet_name == "Head-to-Head":

        append_dict_sheet(
            wb,
            sheet_name,
            paged_rows(
                head_to_head_page,
                db,
                run_id,
            ),
        )

    elif sheet_name == "Missing Categories":

        append_dict_sheet(
            wb,
            sheet_name,
            paged_rows(
                missing_categories_page,
                db,
                run_id,
            ),
        )

    elif sheet_name == "Sitemap Failures":

        append_dict_sheet(
            wb,
            sheet_name,
            failure_rows(
                db,
                run_id,
            ),
        )


def main():

    parser = argparse.ArgumentParser()

    parser.add_argument(
        "--run-id",
        type=int,
        required=True,
    )

    parser.add_argument(
        "--source",
        required=True,
    )

    parser.add_argument(
        "--output",
        required=True,
    )

    args = parser.parse_args()

    source_path = (
        Path(args.source)
        .expanduser()
        .resolve()
    )

    output_path = (
        Path(args.output)
        .expanduser()
        .resolve()
    )

    if not source_path.exists():
        raise SystemExit(
            f"Source workbook not found: "
            f"{source_path}"
        )

    db = SessionLocal()

    source_wb = None

    try:

        run = db.get(
            models.SitemapRun,
            args.run_id,
        )

        if not run:
            raise SystemExit(
                f"Sitemap Run "
                f"{args.run_id} not found"
            )

        print()
        print("==============================")
        print("FULL PARTIAL-RUN EXCEL REPAIR")
        print("==============================")
        print("Run ID :", run.id)
        print("Status :", run.status)
        print("Phase  :", run.phase)
        print("Source :", source_path)
        print("Output :", output_path)
        print()

        source_wb = load_workbook(
            source_path,
            read_only=True,
            data_only=False,
        )

        output_wb = Workbook(
            write_only=True
        )

        existing_sheet_names = (
            source_wb.sheetnames
        )

        for sheet_name in existing_sheet_names:

            if sheet_name in REBUILD_SHEETS:

                rebuild_sheet(
                    output_wb,
                    sheet_name,
                    db,
                    args.run_id,
                )

            else:

                print(
                    f"Copying existing: "
                    f"{sheet_name}"
                )

                copy_existing_sheet(
                    source_wb[sheet_name],
                    output_wb,
                )

        # Defensive handling in case a future/older
        # workbook does not contain one of the
        # rebuildable sheets.
        for sheet_name in REBUILD_SHEETS:

            if (
                sheet_name
                not in existing_sheet_names
            ):

                rebuild_sheet(
                    output_wb,
                    sheet_name,
                    db,
                    args.run_id,
                )

        output_path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        output_wb.save(
            output_path
        )

        print()
        print("==============================")
        print("DONE")
        print("==============================")
        print(output_path)

    finally:

        if source_wb is not None:
            source_wb.close()

        db.close()


if __name__ == "__main__":
    main()