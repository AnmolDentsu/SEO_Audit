from __future__ import annotations

from io import BytesIO
from statistics import median

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app import models
from app.services.benchmarks import benchmark_configuration, build_density_benchmarks, density_target_status
from app.services.content_analysis import build_detailed_analysis, build_content_opportunity_rows


SUCCESS_ROLES = {"CLIENT", "P0", "P1", "OTHER"}


def _append_dict_rows(ws, rows: list[dict], headers: list[str] | None = None):
    rows = rows or []
    if headers is None:
        headers = list(rows[0].keys()) if rows else ["No data"]
    ws.append(headers)
    if not rows:
        ws.append(["No data available for this run."] + [""] * (len(headers) - 1))
        return
    for row in rows:
        ws.append([row.get(h, "") for h in headers])


def _target_priority_order(role: str) -> int:
    return {"CLIENT": 0, "P1": 1, "P0": 2, "OTHER": 3}.get((role or "OTHER").upper(), 9)


def _format_workbook(wb: Workbook):
    for ws in wb.worksheets:
        if ws.max_row > 1:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        sample_rows = min(ws.max_row, 300)
        for col_idx in range(1, ws.max_column + 1):
            max_len = 0
            for row_idx in range(1, sample_rows + 1):
                value = ws.cell(row=row_idx, column=col_idx).value
                if value is not None:
                    max_len = max(max_len, len(str(value)))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 12), 55)

        headers = {cell.column: str(cell.value or "") for cell in ws[1]}
        for col_idx, header in headers.items():
            if "(%)" in header or "Density" in header and "%" in header:
                for row_idx in range(2, ws.max_row + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = '0.00"%"'
        # Long text columns are readable without making the sheet enormous.
        for col_idx, header in headers.items():
            if any(x in header.lower() for x in ["preview", "recommendation", "description", "error"]):
                ws.column_dimensions[get_column_letter(col_idx)].width = 55
                for row_idx in range(2, min(ws.max_row, 1000) + 1):
                    ws.cell(row=row_idx, column=col_idx).alignment = Alignment(wrap_text=True, vertical="top")


def _comparison_rows(keywords, targets, pages, densities, benchmarks):
    page_by_id = {p.id: p for p in pages}
    target_by_id = {t.id: t for t in targets}
    density_by_kw_role = {}
    density_by_kw_label = {}
    for d in densities:
        page = page_by_id.get(d.page_result_id)
        target = target_by_id.get(page.target_id) if page else None
        if not page or not target:
            continue
        role = target.role.upper()
        density_by_kw_role[(d.keyword_id, role)] = d.density_pct
        density_by_kw_label[(d.keyword_id, target.label or target.role)] = d.density_pct

    benchmark_targets = [t for t in targets if t.role.upper() != "CLIENT"]
    raw_labels = [t.label or t.role for t in sorted(targets, key=lambda t: (_target_priority_order(t.role), t.id))]
    rows = []
    for idx, kw in enumerate(keywords, start=1):
        b = benchmarks.get(kw.id, {})
        client_density = density_by_kw_role.get((kw.id, "CLIENT"))
        competitor_values = [
            density_by_kw_label.get((kw.id, t.label or t.role)) for t in benchmark_targets
        ]
        available = [x for x in competitor_values if x is not None]
        row = {
            "Sr. No.": idx,
            "Keyword": kw.keyword,
            "Keyword Type": (kw.keyword_type or "OTHER").upper(),
            "Search Volume": kw.search_volume,
            "Client Source": next((t.label or t.url for t in targets if t.role.upper() == "CLIENT"), ""),
            "Client Density (%)": client_density,
            "Target Min Source": b.get("target_min_source"),
            "Target Min (%)": b.get("target_min"),
            "Target Max Source": b.get("target_max_source"),
            "Target Max (%)": b.get("target_max"),
            "Target Range Assignment": (
                "Unavailable" if b.get("target_min") is None or b.get("target_max") is None
                else "Lower P0/P1 density = Min; higher P0/P1 density = Max"
            ),
            "Benchmark State": b.get("benchmark_state"),
            "Benchmark Ready": "Yes" if b.get("benchmark_ready") else "No",
            "Competitor Pool Count": b.get("competitor_count", 0),
            "Competitor Pool Median Density (%)": b.get("competitor_median"),
            "Client Target Status": density_target_status("CLIENT", client_density, b),
            "Client vs Target Min Gap (%)": None if client_density is None or b.get("target_min") is None else client_density - b["target_min"],
            "Client vs Target Max Gap (%)": None if client_density is None or b.get("target_max") is None else client_density - b["target_max"],
            "Competitor Available Count": len(available),
            "Competitor Expected Count": len(benchmark_targets),
            "Competitor Average Density (%)": sum(available) / len(available) if available else None,
            "Competitor Median Density (%)": median(available) if available else None,
            "Competitor Min Density (%)": min(available) if available else None,
            "Competitor Max Density (%)": max(available) if available else None,
        }
        row["Client vs Competitor Avg Gap (%)"] = None if client_density is None or row["Competitor Average Density (%)"] is None else client_density - row["Competitor Average Density (%)"]
        row["Client vs Competitor Median Gap (%)"] = None if client_density is None or row["Competitor Median Density (%)"] is None else client_density - row["Competitor Median Density (%)"]
        for label in raw_labels:
            target = next(t for t in targets if (t.label or t.role) == label)
            row[f"{label} [{target.role}] Density (%)"] = density_by_kw_label.get((kw.id, label))
        rows.append(row)
    return rows


def _detailed_rows(keywords, targets, pages, densities, benchmarks):
    page_by_id = {p.id: p for p in pages}
    target_by_id = {t.id: t for t in targets}
    keyword_by_id = {k.id: k for k in keywords}
    order = {k.id: i + 1 for i, k in enumerate(keywords)}
    rows = []
    for d in densities:
        page = page_by_id[d.page_result_id]
        target = target_by_id[page.target_id]
        kw = keyword_by_id[d.keyword_id]
        b = benchmarks.get(kw.id, {})
        role = target.role.upper()
        usage = {
            "CLIENT": "Client - compare against the ordered P0/P1 anchor range",
            "P0": "P0 - benchmark anchor; lower/higher role is assigned per keyword",
            "P1": "P1 - benchmark anchor; lower/higher role is assigned per keyword",
        }.get(role, f"{role} - Additional benchmark")
        rows.append({
            "Run ID": page.run_id,
            "Sr. No.": order[kw.id],
            "Keyword": kw.keyword,
            "Keyword Type": (kw.keyword_type or "OTHER").upper(),
            "Search Volume": kw.search_volume,
            "Page Label": target.label or target.role,
            "Priority": role,
            "Role": "Client" if role == "CLIENT" else "Competitor",
            "Benchmark Usage": usage,
            "Requested URL": page.url,
            "Final URL": page.final_url,
            "Keyword Count": d.count,
            "Total Word Count": page.word_count,
            "Density (%)": d.density_pct,
            "Target Min (%)": b.get("target_min"),
            "Target Max (%)": b.get("target_max"),
            "Target Range Assignment": "Unavailable" if b.get("target_min") is None or b.get("target_max") is None else "Lower P0/P1 density = Min; higher P0/P1 density = Max",
            "Benchmark State": b.get("benchmark_state"),
            "Competitor Pool Count": b.get("competitor_count", 0),
            "Competitor Pool Median Density (%)": b.get("competitor_median"),
            "Density Status": density_target_status(role, d.density_pct, b),
            "Fetch Status": "Success" if not page.error else "Unavailable",
            "Evidence ID": f"density:{d.id}",
        })
    rows.sort(key=lambda r: (r["Sr. No."], _target_priority_order(r["Priority"])))
    return rows


def _page_summary_rows(run, targets, pages):
    target_by_id = {t.id: t for t in targets}
    rows = []
    for p in pages:
        t = target_by_id[p.target_id]
        structure = p.structure_json or {}
        rows.append({
            "Run ID": run.id,
            "Fetched At UTC": str(p.fetched_at or ""),
            "Label": t.label or t.role,
            "Priority": t.role,
            "Role": "Client" if t.role.upper() == "CLIENT" else "Competitor",
            "Requested URL": p.url,
            "Final URL": p.final_url,
            "Canonical URL": p.canonical_url,
            "Page Title": p.title,
            "HTTP Status": p.status_code,
            "Word Count": p.word_count,
            "Final Status": "Success" if not p.error else "Failed",
            "Attempt Count": p.retry_count + 1,
            "Success After Retry": "Yes" if not p.error and p.retry_count > 0 else "No",
            "From Cache": "Yes" if p.from_cache else "No",
            "Extraction Root": p.extraction_root or structure.get("root", ""),
            "H1 Text": " | ".join(p.h1_json or []),
            "H2 Count": len(p.h2_json or []),
            "Paragraph Count": p.paragraph_count or structure.get("paragraph_count", 0),
            "FAQ Count": len(structure.get("faqs") or []),
            "Content Hash": p.content_hash,
            "Error": p.error,
            "Evidence ID": f"page:{p.id}",
        })
    return rows


def build_run_workbook(db: Session, run_id: int) -> BytesIO:
    run = db.get(models.Run, run_id)
    if not run:
        raise ValueError("Run not found")
    project = db.get(models.Project, run.project_id)
    targets = db.query(models.Target).filter_by(project_id=project.id).order_by(models.Target.id).all()
    keywords = db.query(models.Keyword).filter_by(project_id=project.id).order_by(models.Keyword.id).all()
    pages = db.query(models.PageResult).filter_by(run_id=run_id).order_by(models.PageResult.id).all()
    densities = (
        db.query(models.DensityResult)
        .join(models.PageResult)
        .filter(models.PageResult.run_id == run_id)
        .order_by(models.DensityResult.id)
        .all()
    )
    opportunities = db.query(models.Opportunity).filter_by(run_id=run_id).order_by(models.Opportunity.id).all()
    target_by_id = {t.id: t for t in targets}
    benchmarks = build_density_benchmarks(pages, densities, target_by_id)
    benchmark_config = benchmark_configuration(targets)

    analysis = dict(run.analysis_json or {})
    if not analysis:
        analysis = build_detailed_analysis(pages, target_by_id, keywords, opportunities)
    elif opportunities:
        analysis["content_opportunities"] = build_content_opportunity_rows(
            analysis.get("section_comparison", []),
            analysis.get("keyword_placement", []),
            analysis.get("faq_comparison", []),
            opportunities,
        )

    wb = Workbook()
    wb.remove(wb.active)

    # 1) Run Summary
    ws = wb.create_sheet("Run Summary")
    successful = sum(1 for p in pages if not p.error)
    summary = [
        {"Metric": "URLs processed", "Value": len(pages)},
        {"Metric": "Successful pages", "Value": successful},
        {"Metric": "Final failures", "Value": len(pages) - successful},
        {"Metric": "URL success rate (%)", "Value": round(successful / len(pages) * 100, 2) if pages else 0},
        {"Metric": "Keywords", "Value": len(keywords)},
        {"Metric": "Benchmark configuration", "Value": benchmark_config.get("state")},
        {"Metric": "CLIENT configured", "Value": "Yes" if benchmark_config.get("client_configured") else "No"},
        {"Metric": "P0 configured", "Value": "Yes" if benchmark_config.get("p0_configured") else "No"},
        {"Metric": "P1 configured", "Value": "Yes" if benchmark_config.get("p1_configured") else "No"},
        {"Metric": "OTHER benchmark URLs", "Value": benchmark_config.get("other_count", 0)},
        {"Metric": "Valid keyword-page checks", "Value": len(densities)},
        {"Metric": "Unavailable keyword-page checks", "Value": len(keywords) * len(pages) - len(densities)},
        {"Metric": "Content sections extracted", "Value": len(analysis.get("content_structure", []))},
        {"Metric": "Section comparison topics", "Value": len(analysis.get("section_comparison", []))},
        {"Metric": "FAQ comparison topics", "Value": len(analysis.get("faq_comparison", []))},
        {"Metric": "Content opportunities", "Value": len(analysis.get("content_opportunities", []))},
    ]
    _append_dict_rows(ws, summary)

    # 2) Failure Breakdown
    ws = wb.create_sheet("Failure Breakdown")
    failures = {}
    for p in pages:
        if p.error:
            key = p.error.split(":", 1)[0] if p.error else "UNKNOWN"
            failures[key] = failures.get(key, 0) + 1
    _append_dict_rows(ws, [{"Failure Class": k, "Count": v} for k, v in failures.items()] or [{"Failure Class": "None", "Count": 0}])

    # 3) Run Metadata
    ws = wb.create_sheet("Run Metadata")
    _append_dict_rows(ws, [{
        "Run ID": run.id,
        "Project": project.name,
        "Client": project.client_name or "",
        "Run Status": run.status,
        "Run Phase": run.phase,
        "Created UTC": str(run.created_at),
        "Started UTC": str(run.started_at or ""),
        "Completed UTC": str(run.completed_at or ""),
        "Target Logic": "P0/P1 are benchmark anchors; lower density -> Target Min and higher density -> Target Max per keyword; CLIENT is evaluated; OTHER contributes distribution context",
        "Benchmark Configuration": benchmark_config.get("state"),
        "CLIENT Configured": benchmark_config.get("client_configured"),
        "P0 Configured": benchmark_config.get("p0_configured"),
        "P1 Configured": benchmark_config.get("p1_configured"),
        "OTHER Count": benchmark_config.get("other_count", 0),
        "AI Enabled": bool((run.config_json or {}).get("enable_ai", True)),
        "URLs": len(targets),
        "Keywords": len(keywords),
    }])

    # 4) Comparison Matrix
    ws = wb.create_sheet("Comparison Matrix")
    _append_dict_rows(ws, _comparison_rows(keywords, targets, pages, densities, benchmarks))

    # 5) Detailed Results
    ws = wb.create_sheet("Detailed Results")
    _append_dict_rows(ws, _detailed_rows(keywords, targets, pages, densities, benchmarks))

    # 6) Page Summary
    ws = wb.create_sheet("Page Summary")
    _append_dict_rows(ws, _page_summary_rows(run, targets, pages))

    # 7-11) granular Part 2 tabs restored from the earlier workflow
    for sheet_name, key in [
        ("Content Structure", "content_structure"),
        ("Section Comparison", "section_comparison"),
        ("Keyword Placement", "keyword_placement"),
        ("FAQ Comparison", "faq_comparison"),
        ("Content Opportunities", "content_opportunities"),
    ]:
        ws = wb.create_sheet(sheet_name)
        _append_dict_rows(ws, analysis.get(key, []))

    _format_workbook(wb)
    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out
